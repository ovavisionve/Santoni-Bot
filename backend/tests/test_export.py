"""
Tests for the export service: CSV, Excel, and PDF generation.
"""

import io
import csv
import pytest

from app.services.export_service import (
    export_to_csv,
    export_to_excel,
    export_to_pdf,
    _extract_tables_from_markdown,
)


# ---------------------------------------------------------------------------
# Test data: markdown content with tables
# ---------------------------------------------------------------------------

MARKDOWN_WITH_TABLE = """## Resumen de Ventas
| Zona | Facturas | Total |
| --- | --- | --- |
| Portuguesa | 150 | 1,500,000.00 |
| Barinas | 120 | 1,200,000.00 |
| Lara | 90 | 900,000.00 |
"""

MARKDOWN_PLAIN_TEXT = """Las ventas del mes de enero fueron excelentes.
El total facturado fue de Bs. 3,600,000.00.
Se superaron las metas en un 15%.
"""

MARKDOWN_MULTIPLE_TABLES = """## Ventas por Zona
| Zona | Total |
| --- | --- |
| Portuguesa | 1,500,000.00 |
| Barinas | 1,200,000.00 |

## Ventas por Vendedor
| Vendedor | Total |
| --- | --- |
| Carlos Matias | 800,000.00 |
| Lenny Silva | 700,000.00 |
"""


# ---------------------------------------------------------------------------
# Table extraction from markdown
# ---------------------------------------------------------------------------

class TestExtractTablesFromMarkdown:
    """Tests for the markdown table parser."""

    def test_extract_single_table(self):
        """Should extract one table with headers, rows, and title."""
        tables = _extract_tables_from_markdown(MARKDOWN_WITH_TABLE)
        assert len(tables) == 1

        table = tables[0]
        assert table["title"] == "Resumen de Ventas"
        assert table["headers"] == ["Zona", "Facturas", "Total"]
        assert len(table["rows"]) == 3
        assert table["rows"][0] == ["Portuguesa", "150", "1,500,000.00"]

    def test_extract_multiple_tables(self):
        """Should extract multiple tables from markdown."""
        tables = _extract_tables_from_markdown(MARKDOWN_MULTIPLE_TABLES)
        assert len(tables) == 2
        assert tables[0]["title"] == "Ventas por Zona"
        assert tables[1]["title"] == "Ventas por Vendedor"

    def test_extract_no_tables(self):
        """Plain text without tables should return empty list."""
        tables = _extract_tables_from_markdown(MARKDOWN_PLAIN_TEXT)
        assert tables == []

    def test_extract_empty_string(self):
        """Empty string should return empty list."""
        tables = _extract_tables_from_markdown("")
        assert tables == []


# ---------------------------------------------------------------------------
# CSV export
# ---------------------------------------------------------------------------

class TestCSVExport:
    """Tests for CSV export functionality."""

    def test_csv_with_table(self):
        """CSV export from markdown table should produce valid CSV."""
        result = export_to_csv(MARKDOWN_WITH_TABLE, agent_used="ventas")

        assert isinstance(result, bytes)
        # Decode BOM-prefixed UTF-8
        text = result.decode("utf-8-sig")
        assert len(text) > 0

        # Parse as CSV
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)

        # Should contain title, empty row, header, and data rows
        # Find the header row with column names
        header_found = False
        data_rows = []
        for row in rows:
            if row == ["Zona", "Facturas", "Total"]:
                header_found = True
                continue
            if header_found and len(row) == 3 and row[0]:
                data_rows.append(row)

        assert header_found
        assert len(data_rows) >= 3

    def test_csv_plain_text(self):
        """CSV export from plain text should produce valid CSV."""
        result = export_to_csv(MARKDOWN_PLAIN_TEXT, agent_used="general")

        assert isinstance(result, bytes)
        text = result.decode("utf-8-sig")
        assert "SantoniBot" in text
        assert "general" in text.lower()

    def test_csv_with_agent_name(self):
        """CSV should include agent name when provided."""
        result = export_to_csv(MARKDOWN_PLAIN_TEXT, agent_used="finanzas")
        text = result.decode("utf-8-sig")
        assert "finanzas" in text

    def test_csv_empty_content(self):
        """Exporting empty content should not raise an error."""
        result = export_to_csv("", agent_used=None)
        assert isinstance(result, bytes)

    def test_csv_multiple_tables(self):
        """CSV with multiple tables should include all data."""
        result = export_to_csv(MARKDOWN_MULTIPLE_TABLES)
        text = result.decode("utf-8-sig")

        # Both tables should have their data
        assert "Portuguesa" in text
        assert "Carlos Matias" in text


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

class TestExcelExport:
    """Tests for Excel export functionality."""

    def test_excel_with_table(self):
        """Excel export with table should produce valid Excel bytes."""
        result = export_to_excel(MARKDOWN_WITH_TABLE, agent_used="ventas")
        assert isinstance(result, bytes)

    def test_excel_plain_text(self):
        """Excel export from plain text should produce valid Excel bytes."""
        result = export_to_excel(MARKDOWN_PLAIN_TEXT, agent_used="general")

        assert isinstance(result, bytes)

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws["A1"].value is not None

    def test_excel_plain_text_includes_header(self):
        """Excel plain-text export should contain the SantoniBot header."""
        result = export_to_excel(MARKDOWN_PLAIN_TEXT, agent_used="ventas")

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active

        assert ws.title == "SantoniBot"
        assert "SantoniBot" in str(ws["A1"].value)
        assert "ventas" in str(ws["A2"].value).lower()

    def test_excel_empty_content(self):
        """Exporting empty content should not raise an error."""
        result = export_to_excel("", agent_used=None)
        assert isinstance(result, bytes)

    def test_excel_plain_text_has_content_rows(self):
        """Excel plain-text export should include actual content lines."""
        result = export_to_excel(MARKDOWN_PLAIN_TEXT, agent_used="general")

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active

        flat = ""
        for row in ws.iter_rows(values_only=True):
            flat += " ".join(str(v) for v in row if v is not None) + " "

        assert "ventas" in flat.lower() or "excelentes" in flat.lower()

    def test_excel_multiple_tables(self):
        """Excel with multiple tables should produce valid Excel bytes."""
        result = export_to_excel(MARKDOWN_MULTIPLE_TABLES)
        assert isinstance(result, bytes)


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------

class TestPDFExport:
    """Tests for PDF export functionality."""

    def test_pdf_with_table(self):
        """PDF export from markdown table should produce valid PDF bytes."""
        result = export_to_pdf(MARKDOWN_WITH_TABLE, agent_used="ventas")

        assert isinstance(result, bytes)
        assert len(result) > 0
        # PDF files start with %PDF
        assert result[:5] == b"%PDF-"

    def test_pdf_plain_text(self):
        """PDF export from plain text should produce valid PDF bytes."""
        result = export_to_pdf(MARKDOWN_PLAIN_TEXT, agent_used="general")

        assert isinstance(result, bytes)
        assert result[:5] == b"%PDF-"

    def test_pdf_empty_content(self):
        """Exporting empty content should not raise an error."""
        result = export_to_pdf("", agent_used=None)
        assert isinstance(result, bytes)
        assert result[:5] == b"%PDF-"

    def test_pdf_size_reasonable(self):
        """PDF should be larger than trivial (has actual content)."""
        result = export_to_pdf(MARKDOWN_WITH_TABLE, agent_used="ventas")
        # A PDF with a table should be more than a few hundred bytes
        assert len(result) > 500

    def test_pdf_multiple_tables(self):
        """PDF with multiple tables should not raise an error."""
        result = export_to_pdf(MARKDOWN_MULTIPLE_TABLES, agent_used="ventas")
        assert isinstance(result, bytes)
        assert result[:5] == b"%PDF-"
        assert len(result) > 500

    def test_pdf_with_special_characters(self):
        """PDF should handle Spanish special characters without error."""
        content = (
            "## Resumen de Producci\u00f3n\n"
            "La producci\u00f3n de arroz fue de 1,200 kg.\n"
            "\u00c1rea: Planta Principal\n"
            "A\u00f1o: 2025\n"
        )
        result = export_to_pdf(content, agent_used="produccion")
        assert isinstance(result, bytes)
        assert result[:5] == b"%PDF-"
