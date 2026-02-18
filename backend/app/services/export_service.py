"""
Export service for generating PDF, Excel, and CSV files from chat responses.
"""

import csv
import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


def _extract_tables_from_markdown(text: str) -> list[dict]:
    """Extract tables from markdown text. Returns list of {title, headers, rows}."""
    tables = []
    lines = text.split("\n")
    i = 0

    while i < len(lines):
        line = lines[i].strip()

        # Check for markdown table
        if line.startswith("|") and i + 1 < len(lines):
            # Look for title above
            title = ""
            if i > 0:
                prev = lines[i - 1].strip()
                if prev.startswith("#"):
                    title = prev.lstrip("#").strip()

            # Parse header
            headers = [c.strip() for c in line.split("|")[1:-1]]

            # Skip separator
            if i + 1 < len(lines) and "---" in lines[i + 1]:
                i += 2
            else:
                i += 1

            # Parse rows
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().split("|")[1:-1]]
                rows.append(cells)
                i += 1

            if headers and rows:
                tables.append({"title": title, "headers": headers, "rows": rows})
            continue

        i += 1

    return tables


def export_to_csv(content: str, agent_used: str | None = None) -> bytes:
    """Export chat response content to CSV."""
    output = io.StringIO()
    writer = csv.writer(output)

    tables = _extract_tables_from_markdown(content)

    if tables:
        for table in tables:
            if table["title"]:
                writer.writerow([table["title"]])
                writer.writerow([])
            writer.writerow(table["headers"])
            for row in table["rows"]:
                writer.writerow(row)
            writer.writerow([])
    else:
        # Just write the text content
        writer.writerow(["SantoniBot - Exportación"])
        writer.writerow([f"Agente: {agent_used or 'General'}"])
        writer.writerow([f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        writer.writerow([])
        for line in content.split("\n"):
            clean = line.strip().lstrip("#").strip()
            if clean:
                writer.writerow([clean])

    return output.getvalue().encode("utf-8-sig")


def export_to_excel(content: str, agent_used: str | None = None) -> bytes:
    """Export chat response content to Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "SantoniBot"

    # Styles
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="E06400", end_color="E06400", fill_type="solid")
    title_font = Font(name="Calibri", bold=True, size=14, color="E06400")
    subtitle_font = Font(name="Calibri", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Header
    ws.merge_cells("A1:F1")
    ws["A1"] = "SantoniBot - Reporte"
    ws["A1"].font = title_font

    ws["A2"] = f"Agente: {agent_used or 'General'}"
    ws["A2"].font = subtitle_font
    ws["A3"] = f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"

    current_row = 5
    tables = _extract_tables_from_markdown(content)

    if tables:
        for table in tables:
            if table["title"]:
                ws.cell(row=current_row, column=1, value=table["title"]).font = subtitle_font
                current_row += 1

            # Headers
            for col, header in enumerate(table["headers"], 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
                cell.border = thin_border

            current_row += 1

            # Rows
            for row in table["rows"]:
                for col, value in enumerate(row, 1):
                    cell = ws.cell(row=current_row, column=col, value=value)
                    cell.border = thin_border
                    # Try to convert numeric values
                    try:
                        num = float(value.replace(",", "").replace("Bs.", "").strip())
                        cell.value = num
                        cell.number_format = "#,##0.00"
                    except (ValueError, AttributeError):
                        pass
                current_row += 1

            current_row += 1

        # Auto-width columns
        for col in ws.columns:
            max_length = 0
            col_letter = None
            for cell in col:
                if hasattr(cell, "column_letter"):
                    col_letter = cell.column_letter
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            if col_letter:
                ws.column_dimensions[col_letter].width = min(max_length + 2, 40)
    else:
        # Text content
        for line in content.split("\n"):
            clean = line.strip().lstrip("#").strip()
            if clean:
                ws.cell(row=current_row, column=1, value=clean)
                current_row += 1

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_to_pdf(content: str, agent_used: str | None = None) -> bytes:
    """Export chat response content to PDF."""
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SantoniTitle",
        parent=styles["Title"],
        fontSize=18,
        textColor=colors.HexColor("#042387"),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "SantoniSubtitle",
        parent=styles["Heading2"],
        fontSize=12,
        textColor=colors.HexColor("#042387"),
        spaceAfter=6,
    )
    normal_style = ParagraphStyle(
        "SantoniNormal",
        parent=styles["Normal"],
        fontSize=10,
        spaceAfter=4,
    )

    elements = []

    # Header
    elements.append(Paragraph("SantoniBot - Reporte", title_style))
    elements.append(
        Paragraph(
            f"Agente: {agent_used or 'General'} | "
            f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            normal_style,
        )
    )
    elements.append(Spacer(1, 12))

    tables = _extract_tables_from_markdown(content)

    if tables:
        for table in tables:
            if table["title"]:
                elements.append(Paragraph(table["title"], subtitle_style))

            # Build table data
            table_data = [table["headers"]] + table["rows"]

            # Create table with style
            t = Table(table_data, repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#042387")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, 0), 9),
                        ("FONTSIZE", (0, 1), (-1, -1), 8),
                        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2FF")]),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            elements.append(t)
            elements.append(Spacer(1, 12))
    else:
        # Plain text
        for line in content.split("\n"):
            clean = line.strip()
            if clean.startswith("##"):
                elements.append(Paragraph(clean.lstrip("#").strip(), subtitle_style))
            elif clean:
                elements.append(Paragraph(clean, normal_style))

    doc.build(elements)
    return output.getvalue()


def export_to_docx(content: str, agent_used: str | None = None) -> bytes:
    """Export chat response content to Word (.docx)."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # Styles
    brand_color = RGBColor(0x04, 0x23, 0x87)

    # Title
    title = doc.add_heading("SantoniBot - Reporte", level=1)
    for run in title.runs:
        run.font.color.rgb = brand_color

    # Subtitle
    sub = doc.add_paragraph()
    sub.add_run(f"Agente: {agent_used or 'General'}").bold = True
    sub.add_run(f"  |  Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph()

    tables = _extract_tables_from_markdown(content)

    if tables:
        for table_data in tables:
            if table_data["title"]:
                h = doc.add_heading(table_data["title"], level=2)
                for run in h.runs:
                    run.font.color.rgb = brand_color

            # Create table
            rows = [table_data["headers"]] + table_data["rows"]
            tbl = doc.add_table(rows=len(rows), cols=len(table_data["headers"]))
            tbl.style = "Light Grid Accent 1"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            for i, row in enumerate(rows):
                for j, cell_val in enumerate(row):
                    cell = tbl.cell(i, j)
                    cell.text = cell_val
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            if i == 0:
                                run.bold = True

            doc.add_paragraph()
    else:
        for line in content.split("\n"):
            clean = line.strip()
            if clean.startswith("##"):
                h = doc.add_heading(clean.lstrip("#").strip(), level=2)
                for run in h.runs:
                    run.font.color.rgb = brand_color
            elif clean.startswith("#"):
                h = doc.add_heading(clean.lstrip("#").strip(), level=1)
                for run in h.runs:
                    run.font.color.rgb = brand_color
            elif clean.startswith("- ") or clean.startswith("* "):
                doc.add_paragraph(clean[2:], style="List Bullet")
            elif clean:
                doc.add_paragraph(clean)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
